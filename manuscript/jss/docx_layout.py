"""Keep the editable reading copies aligned with the article's LaTeX sources."""

import re

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


TABLE_WIDTHS = {
    "tab:r-ecosystem": [16, 23, 37, 14, 10],
    "tab:methods": [27, 73],
    "tab:api": [29, 25, 46],
    "tab:datasets": [50, 25, 25],
    "tab:tuned-hnsw": [28, 16, 18, 16, 22],
    "tab:comprehensive-r": [20, 19, 10, 16, 10, 25],
    "tab:supp-environment": [35, 65],
    "tab:dataset-provenance": [18, 23, 45, 14],
    "tab:reference-key": [25, 12, 15, 12, 24, 12],
    "tab:grid": [32, 68],
    "tab:calibration-missing": [15, 45, 20, 20],
    "tab:calibration-stability": [42, 17, 23, 18],
    "tab:calibration-validation": [17, 32, 14, 20, 17],
    "tab:supp-cpu-paired": [28, 12, 12, 12, 12, 12, 12],
    "tab:supp-tuned-hnsw": [26, 17, 18, 18, 21],
    "tab:supp-comprehensive-r": [20, 12, 10, 14, 10, 10, 24],
    "tab:supp-comprehensive-datasets": [18, 20, 20, 21, 21],
    "tab:supp-cuda-paired": [40, 12, 12, 18, 18],
    "tab:supp-cuda-dataset": [28, 12, 12, 12, 12, 12, 12],
    "tab:supp-auto-routes": [60, 20, 20],
    "tab:auto-ivf-parameters": [32, 8, 20, 20, 20],
    "tab:selector-regret-dataset": [24, 12, 15, 13, 12, 15, 13],
    "tab:supp-domain-holdout": [24, 10, 10, 13, 17, 10, 16],
    "tab:supp-derived-methods": [23, 39, 38],
    "tab:supp-evidence-audit": [60, 23, 17],
    "tab:install-matrix": [20, 29, 51],
}


def _is_escaped(text, position):
    """Return whether the character at ``position`` follows an odd slash run."""
    slashes = 0
    position -= 1
    while position >= 0 and text[position] == "\\":
        slashes += 1
        position -= 1
    return slashes % 2 == 1


def _read_braced(text, position):
    """Read one balanced LaTeX braced argument from ``position``."""
    if position >= len(text) or text[position] != "{":
        raise ValueError("Expected a braced LaTeX argument")
    depth = 0
    start = position + 1
    for index in range(position, len(text)):
        if _is_escaped(text, index):
            continue
        if text[index] == "{":
            depth += 1
        elif text[index] == "}":
            depth -= 1
            if depth == 0:
                return text[start:index], index + 1
    raise ValueError("Unbalanced LaTeX argument")


def expand_latex_multicolumns(source):
    """Expand multicolumn cells into ordinary cells for Pandoc's table reader."""
    marker = r"\multicolumn"
    output = []
    position = 0
    while True:
        found = source.find(marker, position)
        if found < 0:
            output.append(source[position:])
            break
        output.append(source[position:found])
        cursor = found + len(marker)
        while cursor < len(source) and source[cursor].isspace():
            cursor += 1
        span_text, cursor = _read_braced(source, cursor)
        while cursor < len(source) and source[cursor].isspace():
            cursor += 1
        _, cursor = _read_braced(source, cursor)
        while cursor < len(source) and source[cursor].isspace():
            cursor += 1
        content, cursor = _read_braced(source, cursor)
        try:
            span = int(span_text)
        except ValueError as error:
            raise ValueError(f"Invalid multicolumn span: {span_text}") from error
        if span < 1:
            raise ValueError(f"Invalid multicolumn span: {span}")
        output.append(content + " &" * (span - 1))
        position = cursor
    return "".join(output)


def apply_reading_layout(document, source, supplementary=False):
    """Set readable table dimensions and preserve source caption numbering."""
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.9)
        for paragraph in section.header.paragraphs:
            paragraph.text = ("faissR supplementary material" if supplementary
                              else "faissR | Journal of Statistical Software manuscript")
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0, 0, 0)
    for style in document.styles:
        if style.type != 1:
            continue
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style.name in {"Normal", "Body Text", "First Paragraph"}:
            style.font.size = Pt(11)
            style.paragraph_format.line_spacing = 1.05
            style.paragraph_format.space_after = Pt(5)
        elif style.name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
        elif style.name == "Source Code":
            style.font.name = "Courier New"
            style.font.size = Pt(8.5)
            style.paragraph_format.line_spacing = 1
    table_caption = figure_caption = 0
    for paragraph in document.paragraphs:
        kind = paragraph.style.name
        if kind == "Source Code":
            paragraph.paragraph_format.keep_together = True
        if kind not in {"Table Caption", "Image Caption", "Figure Caption"}:
            continue
        if kind == "Table Caption":
            table_caption += 1
            label = f"Table {'S' if supplementary else ''}{table_caption}: "
            following = paragraph._p.getnext()
            while following is not None and following.tag == qn("w:p"):
                if following.xpath(".//w:t"):
                    break
                following = following.getnext()
            if following is not None and following.tag == qn("w:tbl"):
                following.addnext(paragraph._p)
        else:
            figure_caption += 1
            label = f"Figure {'S' if supplementary else ''}{figure_caption}: "
        run = paragraph.add_run(label)
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, run._r)
        paragraph.paragraph_format.keep_with_next = False
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.size = Pt(10)
    blocks = re.findall(r"\\begin\{table\}.*?\\end\{table\}", source, re.S)
    table_labels = [
        re.search(r"\\label\{([^}]+)\}", block).group(1)
        for block in blocks
    ]
    labels = []
    for block, label in zip(blocks, table_labels):
        panel_count = len(re.findall(
            r"\\begin\{(?:longtable|tabularx|tabular)\}",
            block,
        ))
        if panel_count <= 1:
            labels.append(label)
        else:
            labels.extend(
                f"{label}::{panel}" for panel in range(1, panel_count + 1)
            )
    if (len(labels) != len(document.tables)
            or table_caption != len(table_labels)):
        raise ValueError("Word tables or captions do not match the source")
    for label, table in zip(labels, document.tables):
        count = len(table.columns)
        proportions = TABLE_WIDTHS.get(label, [1] * count)
        if len(proportions) != count:
            raise ValueError(f"Column count mismatch for {label}")
        widths = [round(9648 * x / sum(proportions)) for x in proportions]
        table.autofit = False
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(width))
            grid.append(column)
        table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:w"), "9648")
        table._tbl.tblPr.find(qn("w:tblW")).set(qn("w:type"), "dxa")
        keep_table = len(table.rows) <= 12
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                start = cell._tc.grid_offset
                span = cell._tc.grid_span
                cell.width = Inches(sum(widths[start:start + span]) / 1440)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.keep_with_next = (
                        keep_table or row_index == len(table.rows) - 1
                    )
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(
                            8.5 if label == "tab:comprehensive-r" else 9.5
                        )
                        if row_index == 0:
                            run.bold = True
